from openpyxl import Workbook,load_workbook
from app.database.db import connect
async def sample(path):
 wb=Workbook(); ws=wb.active; ws.append(['code','name','description','base_price','reseller_price','vip_price','volume','unit','duration','server','stock','low_stock_threshold','categories']); ws.append(['VPN001','نمونه','توضیح',100000,90000,80000,'50','GB','30 روز','DE',10,2,'عمومی|پیشنهاد ویژه']); wb.save(path)
async def import_products(path):
 wb=load_workbook(path,data_only=True); ws=wb.active; h=[str(c.value).strip() if c.value is not None else '' for c in ws[1]]; ix={x:i for i,x in enumerate(h)}
 if not {'code','name','base_price'}<=set(ix): raise ValueError('code,name,base_price لازم است')
 db=await connect(); ok=bad=0; errors=[]
 for rn,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
  try:
   code=str(row[ix['code']]).strip(); name=str(row[ix['name']]).strip(); price=float(row[ix['base_price']]); stock=int(row[ix['stock']]) if 'stock' in ix and row[ix['stock']] is not None else 0
   await db.execute('INSERT INTO products(code,name,description,base_price,reseller_price,vip_price,volume,unit,duration,server,stock,low_stock_threshold) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(code,name,str(row[ix['description']]) if 'description' in ix else '',price,float(row[ix['reseller_price']]) if 'reseller_price' in ix and row[ix['reseller_price']] is not None else None,float(row[ix['vip_price']]) if 'vip_price' in ix and row[ix['vip_price']] is not None else None,str(row[ix['volume']]) if 'volume' in ix else '',str(row[ix['unit']]) if 'unit' in ix else '',str(row[ix['duration']]) if 'duration' in ix else '',str(row[ix['server']]) if 'server' in ix else '',stock,int(row[ix['low_stock_threshold']]) if 'low_stock_threshold' in ix and row[ix['low_stock_threshold']] is not None else 2)); ok+=1
  except Exception as e: bad+=1; errors.append(f'ردیف {rn}: {e}')
 await db.commit(); await db.close(); return ok,bad,errors
